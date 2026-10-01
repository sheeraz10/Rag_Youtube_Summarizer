import math
import re

from pptx import Presentation  # NOTE: from pptx, not from presentation
from pptx.oxml.ns import qn
from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

# --------------------------------
# Design settings
# --------------------------------

SLIDE_W = Inches(13.333)   # 16:9 widescreen
SLIDE_H = Inches(7.5)

NAVY = RGBColor(0x1F, 0x2A, 0x44)
TEXT = RGBColor(0x2B, 0x2B, 0x2B)
MUTED = RGBColor(0x6B, 0x72, 0x80)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FONT = "Calibri"

BODY_SIZE = 26

# How much text fits on one slide before we continue on a new one
MAX_LINES_PER_SLIDE = 9
CHARS_PER_LINE = 75


# --------------------------------
# Markdown parsing
# --------------------------------

def clean_inline(text):
    """Remove markdown inline markers (**bold**, *italic*, `code`)."""
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"__(.+?)__", r"\1", text)
    text = re.sub(r"(?<!\*)\*(?!\s)(.+?)(?<!\s)\*(?!\*)", r"\1", text)
    text = re.sub(r"`(.+?)`", r"\1", text)
    return text.strip()


def parse_summary(summary):
    """
    Turn the markdown summary into a list of sections:
    [{"title": "Key Points", "items": [{"text": ..., "numbered": bool, "level": int}]}]
    """
    sections = []
    current = None

    for raw in summary.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue

        stripped = line.strip()

        # Headings: "## Summary" is the deck title (skipped),
        # "### X" (or any other heading) starts a new slide section
        heading = re.match(r"^(#{1,6})\s+(.*)", stripped)
        if heading:
            level = len(heading.group(1))
            title = clean_inline(heading.group(2))
            if level <= 2:
                continue
            current = {"title": title, "items": []}
            sections.append(current)
            continue

        if current is None:
            current = {"title": "Overview", "items": []}
            sections.append(current)

        indent = len(line) - len(line.lstrip())
        level = 1 if indent >= 2 else 0

        numbered = re.match(r"^\d+[.)]\s+(.*)", stripped)
        bullet = re.match(r"^[-*•]\s+(.*)", stripped)

        if numbered:
            item = {"text": clean_inline(numbered.group(1)), "numbered": True, "level": level}
        elif bullet:
            item = {"text": clean_inline(bullet.group(1)), "numbered": False, "level": level}
        else:
            item = {"text": clean_inline(stripped), "numbered": False, "level": level}

        if item["text"]:
            current["items"].append(item)

    # Drop empty sections
    return [s for s in sections if s["items"]]


# --------------------------------
# Splitting long sections across slides
# --------------------------------

def estimate_lines(text):
    return max(1, math.ceil(len(text) / CHARS_PER_LINE))


def split_items(items):
    """Group items into evenly balanced chunks that each fit on one slide."""
    line_counts = [estimate_lines(i["text"]) for i in items]
    total = sum(line_counts)
    n_chunks = max(1, math.ceil(total / MAX_LINES_PER_SLIDE))
    target = total / n_chunks

    chunks, chunk, used = [], [], 0
    for item, lines in zip(items, line_counts):
        remaining_chunks = n_chunks - len(chunks)
        would_overflow = used + lines > MAX_LINES_PER_SLIDE
        past_target = used >= target and remaining_chunks > 1
        if chunk and (would_overflow or past_target):
            chunks.append(chunk)
            chunk, used = [], 0
        chunk.append(item)
        used += lines

    if chunk:
        chunks.append(chunk)
    return chunks


# --------------------------------
# Slide helpers
# --------------------------------

def set_background(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def style_run(run, size, color, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_title_slide(prs, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    set_background(slide, NAVY)

    t = slide.shapes.title
    t.left, t.top, t.width, t.height = Inches(1), Inches(2.4), Inches(11.333), Inches(1.6)
    t.text_frame.text = title
    t.text_frame.vertical_anchor = MSO_ANCHOR.BOTTOM
    p = t.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    style_run(p.runs[0], 44, WHITE, bold=True)

    s = slide.placeholders[1]
    s.left, s.top, s.width, s.height = Inches(1), Inches(4.2), Inches(11.333), Inches(0.8)
    s.text_frame.text = subtitle
    p = s.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    style_run(p.runs[0], 20, RGBColor(0xC9, 0xD1, 0xE0))


def set_bullet(paragraph, level, numbered, number=None):
    """Real PowerPoint bullets/numbering with a hanging indent."""
    pPr = paragraph._p.get_or_add_pPr()
    indent_in = 0.55 if level == 0 else 0.35
    left_in = (0.55 if level == 0 else 1.0)
    pPr.set("marL", str(int(Inches(left_in))))
    pPr.set("indent", str(-int(Inches(indent_in))))

    for tag in ("a:buNone", "a:buChar", "a:buAutoNum", "a:buClr"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)

    clr = etree.SubElement(pPr, qn("a:buClr"))
    srgb = etree.SubElement(clr, qn("a:srgbClr"))
    srgb.set("val", "1F2A44")

    if numbered:
        bu = etree.SubElement(pPr, qn("a:buAutoNum"))
        bu.set("type", "arabicPeriod")
        if number and number > 1:
            bu.set("startAt", str(number))
    else:
        bu = etree.SubElement(pPr, qn("a:buChar"))
        bu.set("char", "•" if level == 0 else "–")


def add_bullet_paragraph(tf, item, first):
    """Add one bullet/numbered/term paragraph to a text frame."""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_after = Pt(12)

    text = item["text"]
    size = BODY_SIZE - (2 if item["level"] else 0)

    set_bullet(p, item["level"], item["numbered"], item.get("number"))

    # "Term: explanation" -> bold term
    term_match = re.match(r"^([^:]{1,60}):\s+(.*)$", text)

    if term_match and not item["numbered"]:
        r1 = p.add_run()
        r1.text = term_match.group(1) + ": "
        style_run(r1, size, NAVY, bold=True)
        r2 = p.add_run()
        r2.text = term_match.group(2)
        style_run(r2, size, TEXT)
    else:
        r1 = p.add_run()
        r1.text = text
        style_run(r1, size, TEXT)


def add_content_slide(prs, title, items, page_number):
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # "Title Only"
    set_background(slide, WHITE)

    # Title (real title placeholder, repositioned for 16:9)
    t = slide.shapes.title
    t.left, t.top, t.width, t.height = Inches(0.8), Inches(0.5), Inches(11.733), Inches(1.0)
    t.text_frame.text = title
    t.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    t.text_frame.word_wrap = True
    p = t.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    style_run(p.runs[0], 36, NAVY, bold=True)

    # Body text box
    box = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.733), Inches(4.9))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP

    for i, item in enumerate(items):
        add_bullet_paragraph(tf, item, first=(i == 0))

    # Page number
    num = slide.shapes.add_textbox(Inches(11.9), Inches(6.9), Inches(0.9), Inches(0.4))
    num.text_frame.text = str(page_number)
    p = num.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    style_run(p.runs[0], 12, MUTED)


# --------------------------------
# Main function (same signature as before)
# --------------------------------

def create_presentation(title, summary, file_path="youtube_summary.pptx"):
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    add_title_slide(prs, title, "Generated from YouTube video")

    sections = parse_summary(summary)

    page = 2
    for section in sections:
        chunks = split_items(section["items"])

        # Keep numbering continuous across split slides
        counter = 1
        for item in section["items"]:
            if item["numbered"] and item["level"] == 0:
                item["number"] = counter
                counter += 1
            else:
                item["number"] = None

        for idx, chunk in enumerate(chunks):
            slide_title = section["title"]
            if len(chunks) > 1:
                slide_title += f" ({idx + 1}/{len(chunks)})"
            add_content_slide(prs, slide_title, chunk, page)
            page += 1

    # Fallback: if the summary had no headings/content, still make a slide
    if not sections:
        fallback = [{"text": clean_inline(l), "numbered": False, "level": 0, "number": None}
                    for l in summary.splitlines() if l.strip()]
        for idx, chunk in enumerate(split_items(fallback)):
            add_content_slide(prs, f"Summary ({idx + 1})", chunk, page)
            page += 1

    prs.save(file_path)
    return file_path